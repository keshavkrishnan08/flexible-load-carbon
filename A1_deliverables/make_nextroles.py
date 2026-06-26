"""A1 Week 2 Plan — formatted exactly like A1_Week1_Plan (title, blue Role
headings, bold-lead bullets, no preamble). EVERY bullet is hard, load-bearing
math / theory / experiment work, strictly in A1 scope. Outputs Word .docx and a
Word-style (Helvetica) PDF."""
import os
OUT = "/Users/keshavkrishnan/Claude/ASSIPPaper/A1_deliverables"
os.makedirs(OUT, exist_ok=True)
BLUE = (0x1f, 0x56, 0x99)
TITLE = "A1 Week 2 Plan"

ROLES = [
 ("Role 1 (Optimal Contract and General Carbon Theory)", [
   ("Prove existence, uniqueness, and contract optimality.",
    "Prove expected profit is coercive and, under a primitive condition on F (e.g. a log-concave density), "
    "strictly concave on the interior region, so the capacity optimum is unique; and prove the dispatch "
    "contract d*(x)=min{(K_F−x)_+, K_M} is the unique optimal policy via the strict R>c ordering. This "
    "supplies the negative-definite-Hessian condition from primitives instead of assuming it."),
   ("Promote the shift-ratio condition to a primitive.",
    "Prove that a support-wide monotonicity of g(x)=F_{K_R}/f (or a hazard-rate condition on F) implies the "
    "endogenous-window condition the capacity channel currently relies on; give the implication with a proof."),
   ("Prove the necessary-and-sufficient emission condition for general F.",
    "Under that primitive plus a tail/skew condition, prove a sharp iff for carbon reduction at a finite "
    "interior optimum for any such F, upgrading the endogenous inequality (C) and the existence conjecture to "
    "a theorem in primitives. Exhibit a non-power-law F satisfying all conditions to prove the class non-vacuous."),
   ("Prove single-crossing and the viability threshold for general F.",
    "Extend the renewable-cost window from the power-law family to general F and prove Δ(C_R) is single-crossing, "
    "so the carbon-reduction set is the interval {C_R < C_R†} and the strict fall occurs at a finite interior "
    "K_R* rather than the unbounded corner."),
   ("Prove sharpness and tight magnitude bounds.",
    "Prove the shift-ratio condition is necessary as well as sufficient for K_R* > K_R° (a realized boundary, "
    "not just an algebraic sign), and prove the two-sided bounds on the carbon saving Δ are tight, "
    "characterizing the carbon-reduction region in primitive parameter space."),
 ]),
 ("Role 2 (Solver and Numerical Methods)", [
   ("Build and validate the general solver.",
    "Implement the joint-capacity and dispatch solver for d*(x) and (K_R,K_F,K_M) under any residual-demand "
    "distribution, and validate it against the closed forms to machine precision — the verified backend the "
    "rest of the group computes against."),
   ("Run the carbon-sign sweep at a finite optimum.",
    "Sweep over Role 1's primitive condition on a curated distribution family and show, numerically, that "
    "carbon falls at a finite interior optimum — replacing the scattered Week-1 counts with one principled map."),
   ("Certify the Hessian on the interior path.",
    "Numerically certify the profit Hessian is negative definite along the interior equilibrium path for "
    "Role 1's distribution class, by eigenvalue computation, backstopping existence and uniqueness."),
   ("Derive and implement the analytic Hessian and cofactors.",
    "Derive the full analytic Hessian and the comparative-statics cofactors (replacing finite differences, "
    "which are noisy at the knife-edge) and certify them against the solved equilibria across the class."),
 ]),
 ("Role 3 (Comparative Statics and Robustness)", [
   ("Prove the correlation and price comparative statics.",
    "Prove the net-carbon sign is governed by the demand–supply association g = E[S | x]; that mining raises "
    "optimal renewable capacity K_R* > K_R°; and sign ∂(carbon effect)/∂R, ∂/∂β, and ∂/∂p via the "
    "Hessian-cofactor identities, under general conditions."),
   ("Prove the cross-distribution robust certificate.",
    "Sign the capacity covariation dK*/dF along the shape class U so the least-favorable result holds across F "
    "— making C_R ≤ C_R† ⇒ Δ ≥ 0 a genuine distributionally-robust certificate — or prove the obstruction."),
   ("Construct the reversed-channel counterexample.",
    "Build a realized distribution F with g increasing for which K_R* < K_R°, demonstrating the channel reverses "
    "— the sharpness example Week 1 could not produce — and solve it numerically to confirm."),
   ("Prove the limit-interchange rigorously.",
    "Replace the expected-fossil-to-zero corner step with a dominated-convergence or finite-capacity envelope "
    "argument, so the threshold proof holds without the a.s.-limit shortcut."),
 ]),
 ("Role 4 (Mechanism Design and Regimes)", [
   ("Prove the principal–agent implementation with IR and IC.",
    "Cast the regulator as principal and the utility/miner as agent; write the participation (IR) and "
    "incentive-compatibility (IC) constraints and prove the price-based program implements d* optimally and, "
    "where feasible, uniquely — not merely as the principal's first-best dispatch — stating the price-taking / "
    "exogenous-R conditions under which IR and IC are slack."),
   ("Prove the unified regime theorem.",
    "Combine the C_M coverage axis (no adoption / small-FCL / interior / full coverage) and the C_R boundedness "
    "axis into one theorem, and pin down the small-FCL crossover K_M* = K_R° that Week 1 left open."),
   ("Prove the necessity of Assumption 1.",
    "For each of the four parameter conditions, prove it is necessary for the interior optimum or the economic "
    "sign it guards, and characterize the regime obtained when it is relaxed (full-coverage corner, no "
    "adoption, or a perverse curtailment)."),
   ("Characterize the second-best contract under private information.",
    "When the agent privately observes its availability a(x)=(K_F−x)_+, derive the optimal screening contract "
    "(the second-best d*) and characterize how the resulting information rent changes the emission outcome "
    "versus the full-information benchmark."),
 ]),
]

# ---------------- DOCX (Word, Calibri) ----------------
from docx import Document
from docx.shared import Pt, RGBColor, Inches
d = Document()
for s in d.sections:
    s.top_margin=Inches(1.0); s.bottom_margin=Inches(1.0); s.left_margin=s.right_margin=Inches(1.0)
st=d.styles['Normal']; st.font.name='Calibri'; st.font.size=Pt(11)
tt=d.add_paragraph(); rt=tt.add_run(TITLE); rt.bold=True; rt.font.size=Pt(20); rt.font.color.rgb=RGBColor(*BLUE)
for head, bullets in ROLES:
    h=d.add_paragraph(); h.paragraph_format.space_before=Pt(10)
    rh=h.add_run(head); rh.bold=True; rh.font.size=Pt(13); rh.font.color.rgb=RGBColor(*BLUE)
    for lead, body in bullets:
        bp=d.add_paragraph(style='List Bullet'); bl=bp.add_run(lead+" "); bl.bold=True; bp.add_run(body)
docx_path=os.path.join(OUT,"A1_Week2_NextRoles.docx"); d.save(docx_path); print("wrote",docx_path)

# ---------------- PDF (Word-style, Helvetica) ----------------
from reportlab.lib.pagesizes import letter
from reportlab.lib.units import inch
from reportlab.lib.colors import Color
from reportlab.platypus import SimpleDocTemplate, Paragraph, ListFlowable, ListItem
from reportlab.lib.styles import ParagraphStyle
blue=Color(BLUE[0]/255,BLUE[1]/255,BLUE[2]/255)
H1=ParagraphStyle('H1',fontName='Helvetica-Bold',fontSize=20,textColor=blue,spaceAfter=10)
RH=ParagraphStyle('RH',fontName='Helvetica-Bold',fontSize=13,textColor=blue,spaceBefore=12,spaceAfter=4)
BUL=ParagraphStyle('BUL',fontName='Helvetica',fontSize=11,leading=14)
def esc(s): return s.replace("&","&amp;")
pdf_path=os.path.join(OUT,"A1_Week2_NextRoles.pdf")
doc=SimpleDocTemplate(pdf_path,pagesize=letter,topMargin=1*inch,bottomMargin=1*inch,leftMargin=1*inch,rightMargin=1*inch)
fl=[Paragraph(TITLE,H1)]
for head,bullets in ROLES:
    fl.append(Paragraph(esc(head),RH))
    items=[ListItem(Paragraph(f"<b>{esc(lead)}</b> {esc(body)}",BUL),leftIndent=14,spaceAfter=5) for lead,body in bullets]
    fl.append(ListFlowable(items,bulletType='bullet',start='•',leftIndent=12))
doc.build(fl); print("wrote",pdf_path)
